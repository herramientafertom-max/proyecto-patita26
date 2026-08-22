#!/usr/bin/env python3
from collections import Counter, defaultdict
from pathlib import Path
import re

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DB_FILE = ROOT / "datos-origen" / "Base_de_Datos_1.xlsx"
CONTROL_FILE = ROOT / "datos-origen" / "Clientes_IDs_Maestra_de_Locales_y_Registro_de_Errores.xlsx"


def txt(value):
    if value is None:
        return ""
    return str(value).strip()


def norm(value):
    return re.sub(r"\s+", " ", txt(value)).casefold()


def rows(ws):
    return ws.iter_rows(min_row=2, values_only=True)


db = load_workbook(DB_FILE, read_only=True, data_only=True)
control = load_workbook(CONTROL_FILE, read_only=True, data_only=True)

locales = list(rows(control["Maestra de Locales"]))
local_by_code = {txt(r[0]).upper().replace(" ", ""): r for r in locales if txt(r[0])}
local_names = defaultdict(list)
for r in locales:
    if txt(r[4]):
        local_names[norm(r[4])].append(r)

products = list(rows(db["Maestra Productos"]))
product_keys = {txt(r[0]).upper().replace(" ", ""): r for r in products if txt(r[0])}
product_sap = defaultdict(list)
for r in products:
    if txt(r[2]):
        product_sap[txt(r[2])].append(r)

prefixes = Counter(re.match(r"^[A-Za-z]+", txt(r[0])).group(0).upper() if re.match(r"^[A-Za-z]+", txt(r[0])) else "SIN_PREFIJO" for r in products)

def resolution(sheet_name, local_col, product_col, local_is_code=False):
    total = local_missing = local_ambiguous = product_missing = product_ambiguous = 0
    for r in rows(db[sheet_name]):
        if not any(v is not None and txt(v) for v in r):
            continue
        total += 1
        local_value = txt(r[local_col])
        if local_is_code:
            matches = [local_by_code.get(local_value.upper().replace(" ", ""))]
            matches = [x for x in matches if x]
        else:
            matches = local_names.get(norm(local_value), [])
        if not matches:
            local_missing += 1
            continue
        if len(matches) > 1:
            local_ambiguous += 1
        chain = txt(matches[0][2]).upper()
        sap_or_key = txt(r[product_col]).upper().replace(" ", "")
        if sap_or_key in product_keys:
            pmatches = [product_keys[sap_or_key]]
        else:
            pmatches = [p for p in product_sap.get(txt(r[product_col]), []) if chain[:3] in txt(p[0]).upper()]
            if not pmatches:
                pmatches = product_sap.get(txt(r[product_col]), [])
        if not pmatches:
            product_missing += 1
        elif len(pmatches) > 1:
            product_ambiguous += 1
    print(f"{sheet_name}: total={total} local_missing={local_missing} local_ambiguous={local_ambiguous} product_missing={product_missing} product_ambiguous={product_ambiguous}")

print("product_prefixes=" + ",".join(f"{k}:{v}" for k, v in sorted(prefixes.items())))
print(f"products={len(products)} unique_product_keys={len(product_keys)} unique_sap={len(product_sap)} duplicate_sap_values={sum(1 for v in product_sap.values() if len(v)>1)}")
print(f"locals={len(locales)} unique_local_codes={len(local_by_code)} duplicate_local_names={sum(1 for v in local_names.values() if len(v)>1)}")

resolution("Maestra Retail", 1, 2, local_is_code=True)
resolution("Historial de Stock", 1, 1, local_is_code=True)
resolution("Inventario y Sugerido", 5, 7)
resolution("Fotos en Terreno", 6, 10)

users = list(rows(db["Usuarios"]))
names_to_ruts = defaultdict(set)
for r in users:
    rut = re.sub(r"\D", "", txt(r[0]))
    if rut and txt(r[2]):
        names_to_ruts[norm(r[2])].add(rut)
access_total = access_missing = access_ambiguous = 0
for r in rows(db["Registro de Accesos"]):
    if not any(v is not None and txt(v) for v in r):
        continue
    access_total += 1
    matches = names_to_ruts.get(norm(r[1]), set())
    access_missing += not matches
    access_ambiguous += len(matches) > 1
print(f"Registro de Accesos: total={access_total} user_missing={access_missing} user_ambiguous={access_ambiguous}")

geo = Counter()
for r in locales:
    value = txt(r[5])
    if not value:
        geo["empty"] += 1
    elif re.search(r"-?\d+[.,]\d+\s*[,;]\s*-?\d+[.,]\d+", value):
        geo["coordinate_pair"] += 1
    elif "maps" in value.casefold() or "http" in value.casefold():
        geo["url"] += 1
    else:
        geo["other"] += 1
print("geolocation=" + ",".join(f"{k}:{v}" for k, v in sorted(geo.items())))

error_rows = [r for r in rows(control["Registro de Errores"]) if any(v is not None and txt(v) for v in r)]
print(f"Registro de Errores: nonempty_rows={len(error_rows)}")
